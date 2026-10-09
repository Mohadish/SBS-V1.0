# -*- coding: utf-8 -*-
"""DEMO.pdf → DEMO-steps.xlsx for SBS "Steps from Excel".
Figures are RENDERED from the page at 300 dpi (Print-To-PDF sliced every screenshot into strips; rendering the
figure region recomposes them at their native pixel density — no re-encoding of a bitmap).
The xlsx keeps the full-resolution PNG bytes; the picture is only DISPLAYED smaller in Excel."""
import os, re, json, sys
import fitz
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
sys.path.insert(0, os.path.dirname(__file__))
from steps_data import ROWS, CH

PDF = r'D:\Nadav-Avi\project\Elbit\doc-demo\DEMO.pdf'
OUTDIR = r'D:\Nadav-Avi\project\Elbit\doc-demo'
IMGDIR = os.path.join(OUTDIR, 'DEMO-images')
XLSX = os.path.join(OUTDIR, 'DEMO-steps.xlsx')
DPI = 300
os.makedirs(IMGDIR, exist_ok=True)
doc = fitz.open(PDF)
struct = {r['page']: r for r in json.load(open(os.path.join(os.path.dirname(__file__), 'structure.json')))}

# ── figure regions: each image cluster belongs to the first "Figure N:" caption below it on its page ──
regions = {}       # 'F23' → (page_index, Rect)
captions = {}      # 'F23' → caption text
for pg, r in struct.items():
    caps = sorted(r['captions'], key=lambda c: c['y'])
    prev_y = 40
    for c in caps:
        m = re.match(r'^\s*Figure (\d+)\s*:\s*(.*)$', c['text'])
        if not m: continue
        key = 'F' + m.group(1)
        if key in captions: prev_y = c['y']; continue           # the TOC lists every figure too — first real one wins (TOC pages have no clusters)
        cl = [fitz.Rect(*b) for b in r['img_clusters'] if b[3] <= c['y'] + 1 and b[1] >= prev_y - 1]
        if cl:
            big = [b for b in cl if b.height >= 25] or cl
            u = fitz.Rect(big[0])
            for b in big: u.include_rect(b)
            regions[key] = (pg - 1, u)
            captions[key] = c['text'].strip()
        prev_y = c['y']
# vector / special regions (hand-measured)
regions['F3'] = (8, fitz.Rect(23, 178, 571, 516)); captions['F3'] = 'Figure 3: Production Flow'
regions['T1'] = (9, fitz.Rect(55, 143, 554, 283)); captions['T1'] = 'Table 1: Production Files'
regions['T2'] = (10, fitz.Rect(47, 106, 553, 318)); captions['T2'] = 'Table 2: Customer SW Files'
regions['F22'] = (19, fitz.Rect(40, 94, 556, 290))          # both configurations with their labels
regions['F36'] = (25, fitz.Rect(169, 114, 444, 233))        # without the inline button icon of step 11
regions['F42'] = (28, fitz.Rect(139, 230, 456, 478))        # without the inline FTP icon of step 2
regions['F53'] = (34, fitz.Rect(132, 163, 484, 389))        # without the icon above
missing = [k for k in range(1, 62) if 'F%d' % k not in regions]
print('figure regions:', len(regions), '| missing:', missing)

# ── render ──
Z = fitz.Matrix(DPI / 72, DPI / 72)
images = {}
for key, (pi, rect) in sorted(regions.items(), key=lambda kv: (kv[1][0], kv[1][1].y0)):
    page = doc[pi]
    clip = (rect + (-3, -3, 3, 3)) & page.rect
    pix = page.get_pixmap(matrix=Z, clip=clip, alpha=False)
    path = os.path.join(IMGDIR, '%s.png' % key.replace('F', 'fig').replace('T', 'table'))
    pix.save(path)
    images[key] = (path, pix.width, pix.height)
print('rendered', len(images), 'images to', IMGDIR)

# ── workbook ──
wb = Workbook()
ws = wb.active; ws.title = 'Steps'
HEAD = ['Step', 'Voiceover', 'Notes', 'Image', 'Image 2', 'Title', 'Chapter', 'Source (PDF page / step)']
ws.append(HEAD)
widths = [30, 72, 48, 46, 46, 34, 42, 18]
for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='2F5597'); c.alignment = Alignment(vertical='center')
ws.freeze_panes = 'A2'
DISPLAY_W = 320          # px shown in Excel; the stored PNG stays full size
chapter_rows = {}
for i, r in enumerate(ROWS, start=2):
    ws.cell(i, 1, r['name']); ws.cell(i, 2, r['voice']); ws.cell(i, 3, r['notes'] or None)
    ws.cell(i, 6, r['title'] or None); ws.cell(i, 7, '%s %s' % (r['ch'], CH[r['ch']])); ws.cell(i, 8, 'p%s / %s' % (r['page'], r['docstep']))
    for col in (1, 2, 3, 6, 7, 8): ws.cell(i, col).alignment = Alignment(wrap_text=True, vertical='top')
    chapter_rows.setdefault(r['ch'], []).append(i)
    hmax = 48
    for k, fig in enumerate(r['figs'][:2]):
        if fig not in images: print('  !! row', i, 'missing figure', fig); continue
        path, w, h = images[fig]
        img = XLImage(path)
        sc = DISPLAY_W / w
        img.width, img.height = int(w * sc), int(h * sc)
        ws.add_image(img, '%s%d' % (get_column_letter(4 + k), i))
        hmax = max(hmax, img.height * 0.75 + 6)
    ws.row_dimensions[i].height = min(hmax, 400)
# import plan sheet: the row-selection line for the dialog, chapters in document order
def compress(rows):
    out = []; s = p = rows[0]
    for x in rows[1:]:
        if x == p + 1: p = x; continue
        out.append('%d-%d' % (s, p) if s != p else str(s)); s = p = x
    out.append('%d-%d' % (s, p) if s != p else str(s))
    return ','.join(out)
line = ', '.join('%s(%s)%s' % (code, CH[code], compress(rows)) for code, rows in chapter_rows.items())
ws2 = wb.create_sheet('Import plan')
ws2['A1'] = 'How to import into SBS'; ws2['A1'].font = Font(bold=True, size=13)
ws2['A2'] = 'Files ▸ Steps from Excel… ▸ pick DEMO-steps.xlsx ▸ sheet "Steps" ▸ tick "first row is a header".'
ws2['A3'] = 'Roles: A = Step name, B = Voiceover, D and E = Image, F = Title (optional), C/G/H = Ignore.'
ws2['A4'] = 'Paste this line into "Rows to import" to get the chapters (Excel row numbers, header = row 1):'
ws2['A5'] = line; ws2['A5'].alignment = Alignment(wrap_text=True, vertical='top'); ws2.row_dimensions[5].height = 90
ws2['A7'] = 'Source: DEMO.pdf (020389D-06 Rev A-), %d pages, %d figures rendered at %d dpi to DEMO-images\\.' % (len(doc), len(images), DPI)
ws2['A8'] = 'Known source defects: step 4.1-18 has a broken cross-reference in the PDF (taken as "step 23"); section 7 numbers its last step "12".'
ws2.column_dimensions['A'].width = 140
wb.save(XLSX)
print('wrote', XLSX, '| rows', len(ROWS), '| chapters', len(chapter_rows))
print('row line:', line)
json.dump({k: {'path': v[0], 'w': v[1], 'h': v[2], 'caption': captions.get(k, '')} for k, v in images.items()}, open(os.path.join(os.path.dirname(__file__), 'images.json'), 'w'), indent=1)
